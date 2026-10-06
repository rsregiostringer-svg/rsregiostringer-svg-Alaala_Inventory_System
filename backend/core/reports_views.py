import csv
import io
from django.http import HttpResponse
from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework import permissions
from django.db.models import Sum, Count, Q
from inventory.models import InventoryItem, InventoryTransaction
from laundry.models import LaundryRecord
from caskets.models import Casket, CasketSale, LamayRecord, Chapel
from maintenance.models import Maintenance
from utilities.models import WaterBill, ElectricityBill
from .permissions import CanViewReports


class ReportsSummaryView(APIView):
    permission_classes = [CanViewReports]

    def get(self, request):
        user = request.user
        accessible_loc_ids = user.get_accessible_location_ids()
        location_id = request.query_params.get('location')
        start_date = request.query_params.get('start_date')
        end_date = request.query_params.get('end_date')

        # Filter helpers with strict location isolation
        def apply_loc(qs, field='location_id'):
            if not user.is_master_admin:
                qs = qs.filter(**{f"{field}__in": accessible_loc_ids})
            if location_id:
                try:
                    loc_int = int(location_id)
                    if user.is_master_admin or loc_int in accessible_loc_ids:
                        return qs.filter(**{field: loc_int})
                    else:
                        return qs.none()
                except ValueError:
                    return qs.filter(**{field: location_id})
            return qs

        # Inventory Stats
        items_qs = apply_loc(InventoryItem.objects.filter(is_active=True))
        total_items = items_qs.count()
        total_qty = items_qs.aggregate(Sum('current_quantity'))['current_quantity__sum'] or 0
        total_cost_val = sum(item.current_quantity * float(item.cost) for item in items_qs)
        low_stock = items_qs.filter(status='LOW_STOCK').count()
        out_of_stock = items_qs.filter(status='OUT_OF_STOCK').count()

        # Laundry Stats
        laundry_qs = apply_loc(LaundryRecord.objects.all())
        if start_date:
            laundry_qs = laundry_qs.filter(laundry_in_date__gte=start_date)
        if end_date:
            laundry_qs = laundry_qs.filter(laundry_in_date__lte=end_date)
        laundry_total = laundry_qs.count()
        laundry_in_process = laundry_qs.exclude(status='RETURNED').count()
        laundry_returned = laundry_qs.filter(status='RETURNED').count()

        # Caskets Stats
        caskets_qs = apply_loc(Casket.objects.all())
        caskets_available = caskets_qs.filter(status='AVAILABLE').aggregate(Sum('quantity'))['quantity__sum'] or 0
        caskets_reserved = caskets_qs.filter(status='RESERVED').aggregate(Sum('quantity'))['quantity__sum'] or 0
        caskets_sold = caskets_qs.filter(status='SOLD').aggregate(Sum('quantity'))['quantity__sum'] or 0
        caskets_repair = caskets_qs.filter(status='FOR_REPAIR').aggregate(Sum('quantity'))['quantity__sum'] or 0

        casket_sales_qs = CasketSale.objects.all()
        # Ensure we only calculate sales for caskets that belong to accessible locations
        if not user.is_master_admin:
            casket_sales_qs = casket_sales_qs.filter(casket__location_id__in=accessible_loc_ids)
        if location_id:
            casket_sales_qs = casket_sales_qs.filter(casket__location_id=location_id)
        if start_date:
            casket_sales_qs = casket_sales_qs.filter(date_sold__gte=start_date)
        if end_date:
            casket_sales_qs = casket_sales_qs.filter(date_sold__lte=end_date)
        
        caskets_revenue = casket_sales_qs.aggregate(Sum('selling_price'))['selling_price__sum'] or 0

        # Maintenance Stats
        maint_qs = apply_loc(Maintenance.objects.all())
        if start_date:
            maint_qs = maint_qs.filter(date_reported__gte=start_date)
        if end_date:
            maint_qs = maint_qs.filter(date_reported__lte=end_date)
        maint_total = maint_qs.count()
        maint_completed = maint_qs.filter(status='COMPLETED').count()
        maint_open = maint_qs.filter(status__in=['REPORTED', 'PENDING', 'FOR_REPAIR', 'IN_PROGRESS']).count()
        maint_cost = maint_qs.aggregate(Sum('cost'))['cost__sum'] or 0

        # Utilities Stats
        water_qs = apply_loc(WaterBill.objects.all())
        elec_qs = apply_loc(ElectricityBill.objects.all())
        if start_date:
            water_qs = water_qs.filter(due_date__gte=start_date)
            elec_qs = elec_qs.filter(due_date__gte=start_date)
        if end_date:
            water_qs = water_qs.filter(due_date__lte=end_date)
            elec_qs = elec_qs.filter(due_date__lte=end_date)

        water_amount = water_qs.aggregate(Sum('amount'))['amount__sum'] or 0
        water_consumption = water_qs.aggregate(Sum('consumption'))['consumption__sum'] or 0
        water_unpaid = water_qs.filter(payment_status__in=['UNPAID', 'OVERDUE']).aggregate(Sum('amount'))['amount__sum'] or 0

        elec_amount = elec_qs.aggregate(Sum('amount'))['amount__sum'] or 0
        elec_consumption = elec_qs.aggregate(Sum('consumption'))['consumption__sum'] or 0
        elec_unpaid = elec_qs.filter(payment_status__in=['UNPAID', 'OVERDUE']).aggregate(Sum('amount'))['amount__sum'] or 0

        return Response({
            'inventory': {
                'total_items': total_items,
                'total_quantity': total_qty,
                'valuation': float(total_cost_val),
                'low_stock': low_stock,
                'out_of_stock': out_of_stock,
            },
            'laundry': {
                'total': laundry_total,
                'in_process': laundry_in_process,
                'returned': laundry_returned,
            },
            'caskets': {
                'available': caskets_available,
                'reserved': caskets_reserved,
                'sold': caskets_sold,
                'for_repair': caskets_repair,
                'revenue': float(caskets_revenue),
            },
            'chapels': {
                'total': Chapel.objects.filter(is_active=True).count(),
                'available': Chapel.objects.filter(status=Chapel.Status.AVAILABLE, is_active=True).count(),
                'occupied': Chapel.objects.filter(status=Chapel.Status.OCCUPIED, is_active=True).count(),
                'cleaning': Chapel.objects.filter(status=Chapel.Status.CLEANING, is_active=True).count(),
                'active_lamay': LamayRecord.objects.filter(status=LamayRecord.Status.ACTIVE).count(),
            },
            'maintenance': {
                'total_tickets': maint_total,
                'open_tickets': maint_open,
                'completed_tickets': maint_completed,
                'total_cost': float(maint_cost),
            },
            'utilities': {
                'water_amount': float(water_amount),
                'water_consumption': float(water_consumption),
                'water_unpaid': float(water_unpaid),
                'electricity_amount': float(elec_amount),
                'electricity_consumption': float(elec_consumption),
                'electricity_unpaid': float(elec_unpaid),
                'total_utility_cost': float(water_amount + elec_amount),
                'total_unpaid': float(water_unpaid + elec_unpaid),
            }
        })


class ExportCSVView(APIView):
    permission_classes = [CanViewReports]

    def get(self, request):
        user = request.user
        accessible_loc_ids = user.get_accessible_location_ids()
        module = request.query_params.get('module', 'inventory').lower()
        location_id = request.query_params.get('location')
        start_date = request.query_params.get('start_date')
        end_date = request.query_params.get('end_date')

        filename = f"alaala_{module}_report"
        if module == 'laundry':
            if start_date and end_date:
                filename = f"alaala_laundry_{start_date}_to_{end_date}"
            elif start_date:
                filename = f"alaala_laundry_from_{start_date}"
            elif end_date:
                filename = f"alaala_laundry_until_{end_date}"
            else:
                filename = "alaala_laundry_all_records"

        response = HttpResponse(content_type='text/csv')
        response['Content-Disposition'] = f'attachment; filename="{filename}.csv"'
        writer = csv.writer(response)

        # Helper to check location authorization
        def get_target_loc_filter(qs, loc_field='location_id'):
            if not user.is_master_admin:
                qs = qs.filter(**{f"{loc_field}__in": accessible_loc_ids})
            if location_id:
                try:
                    loc_int = int(location_id)
                    if user.is_master_admin or loc_int in accessible_loc_ids:
                        qs = qs.filter(**{loc_field: loc_int})
                    else:
                        return qs.none()
                except ValueError:
                    qs = qs.filter(**{loc_field: location_id})
            return qs

        if module == 'inventory':
            writer.writerow(['ID', 'Item Name', 'Category', 'Location', 'Quantity', 'Unit', 'Min Stock', 'Cost (PHP)', 'Total Value (PHP)', 'Status'])
            items = InventoryItem.objects.select_related('category', 'location').filter(is_active=True)
            items = get_target_loc_filter(items, 'location_id')
            for it in items:
                val = it.current_quantity * float(it.cost)
                writer.writerow([it.id, it.item_name, it.category.name, it.location.name, it.current_quantity, it.unit, it.minimum_stock, it.cost, f"{val:.2f}", it.status])

        elif module == 'transactions':
            writer.writerow(['Date', 'Type', 'Item', 'Qty', 'From Location', 'To Location', 'Reason', 'User', 'Notes'])
            txs = InventoryTransaction.objects.select_related('item', 'from_location', 'to_location', 'user').all()
            if not user.is_master_admin:
                txs = txs.filter(Q(from_location_id__in=accessible_loc_ids) | Q(to_location_id__in=accessible_loc_ids))
            if location_id:
                txs = txs.filter(Q(from_location_id=location_id) | Q(to_location_id=location_id))
            if start_date and start_date.strip():
                txs = txs.filter(created_at__date__gte=start_date.strip())
            if end_date and end_date.strip():
                txs = txs.filter(created_at__date__lte=end_date.strip())
            for t in txs:
                writer.writerow([
                    t.created_at.strftime('%Y-%m-%d %H:%M'),
                    t.transaction_type,
                    t.item.item_name,
                    t.quantity,
                    t.from_location.name if t.from_location else '',
                    t.to_location.name if t.to_location else '',
                    t.reason,
                    t.user.username if t.user else '',
                    t.notes
                ])

        elif module == 'laundry':
            writer.writerow([
                'Laundry IN (Date/ Shift/ In Charge)',
                'Items',
                'Quantity',
                'Laba (Date/Shift/ In Charge)',
                'Banlaw (Date/Shift/ In Charge)',
                'Sampay (Date/Shift/ In Charge)',
                'Pinaw (Date/Shift/ In Charge)',
                'Tiklop (Date/Shift/ In Charge)',
                'Date Returned/ By',
                'Encoded By',
                'Status'
            ])
            recs = LaundryRecord.objects.select_related('location').all()
            recs = get_target_loc_filter(recs, 'location_id')
            if start_date and start_date.strip():
                recs = recs.filter(laundry_in_date__gte=start_date.strip())
            if end_date and end_date.strip():
                recs = recs.filter(laundry_in_date__lte=end_date.strip())
            for r in recs:
                laundry_in_dt = f"{r.laundry_in_date} {r.laundry_in_time}".strip()
                laundry_in_str = f"{laundry_in_dt} / {r.laundry_in_shift} / {r.laundry_in_charge}"

                laba_dt = f"{r.laba_date} {r.laba_time}".strip() if r.laba_date else ""
                laba_str = f"{laba_dt} / {r.laba_shift or ''} / {r.laba_in_charge or ''}" if r.laba_date else ""

                banlaw_dt = f"{r.banlaw_date} {r.banlaw_time}".strip() if r.banlaw_date else ""
                banlaw_str = f"{banlaw_dt} / {r.banlaw_shift or ''} / {r.banlaw_in_charge or ''}" if r.banlaw_date else ""

                sampay_dt = f"{r.sampay_date} {r.sampay_time}".strip() if r.sampay_date else ""
                sampay_str = f"{sampay_dt} / {r.sampay_shift or ''} / {r.sampay_in_charge or ''}" if r.sampay_date else ""

                pinaw_dt = f"{r.pinaw_date} {r.pinaw_time}".strip() if r.pinaw_date else ""
                pinaw_str = f"{pinaw_dt} / {r.pinaw_shift or ''} / {r.pinaw_in_charge or ''}" if r.pinaw_date else ""

                tiklop_dt = f"{r.tiklop_date} {r.tiklop_time}".strip() if r.tiklop_date else ""
                tiklop_str = f"{tiklop_dt} / {r.tiklop_shift or ''} / {r.tiklop_in_charge or ''}" if r.tiklop_date else ""

                returned_dt = f"{r.date_returned} {r.returned_time}".strip() if r.date_returned else ""
                returned_str = f"{returned_dt} / {r.returned_by or ''}" if r.date_returned else ""

                writer.writerow([
                    laundry_in_str,
                    f"{r.item} ({r.location.name})",
                    r.quantity,
                    laba_str,
                    banlaw_str,
                    sampay_str,
                    pinaw_str,
                    tiklop_str,
                    returned_str,
                    r.encoded_by,
                    r.get_status_display()
                ])

        elif module == 'caskets':
            writer.writerow(['Casket ID', 'Model', 'Type', 'Size', 'Color', 'Material', 'Location', 'Cost (PHP)', 'Price (PHP)', 'Condition', 'Status', 'Received Date'])
            caskets = Casket.objects.select_related('location').all()
            caskets = get_target_loc_filter(caskets, 'location_id')
            for c in caskets:
                writer.writerow([
                    c.casket_id, c.model, c.casket_type, c.size, c.color, c.material,
                    c.location.name, c.purchase_cost, c.selling_price, c.condition, c.status, c.date_received
                ])

        elif module == 'casket_sales':
            # Requirement 24: CASKET SALES REPORT
            writer.writerow(['Casket ID', 'Model', 'Buyer', 'Buyer Contact', 'Deceased', 'Selling Price (PHP)', 'Date Sold', 'Chapel', 'Encoded By'])
            sales = CasketSale.objects.select_related('casket', 'buyer', 'deceased', 'chapel', 'encoded_by').all()
            if start_date:
                sales = sales.filter(date_sold__date__gte=start_date)
            if end_date:
                sales = sales.filter(date_sold__date__lte=end_date)
            chapel_id = request.query_params.get('chapel')
            if chapel_id:
                sales = sales.filter(chapel_id=chapel_id)
            for s in sales:
                writer.writerow([
                    s.casket.casket_id,
                    s.casket.model,
                    s.buyer.full_name,
                    s.buyer.contact_number,
                    s.deceased.full_name,
                    f"{s.selling_price:.2f}",
                    s.date_sold.strftime('%Y-%m-%d %H:%M'),
                    s.chapel.name if s.chapel else (f"Residence: {s.residence_address}" if s.residence_address else 'Residence'),
                    s.encoded_by.username if s.encoded_by else 'System'
                ])

        elif module == 'lamay':
            # Requirement 24: LAMAY REPORT
            writer.writerow(['Lamay ID', 'Chapel', 'Deceased', 'Buyer/Family Contact', 'Casket', 'Lamay Start', 'Expected End', 'Burial Date', 'Status'])
            lamays = LamayRecord.objects.select_related('chapel', 'deceased', 'buyer', 'casket').all()
            if start_date:
                lamays = lamays.filter(lamay_start_date__gte=start_date)
            if end_date:
                lamays = lamays.filter(lamay_start_date__lte=end_date)
            chapel_id = request.query_params.get('chapel')
            if chapel_id:
                lamays = lamays.filter(chapel_id=chapel_id)
            status_param = request.query_params.get('status')
            if status_param:
                lamays = lamays.filter(status=status_param)
            for l in lamays:
                burial_str = f"{l.expected_burial_date} {l.burial_time or ''}".strip() if l.expected_burial_date else 'TBD'
                venue_str = l.chapel.name if l.chapel else (f"Residence: {l.residence_address}" if l.residence_address else 'Residence / Home Viewing')
                writer.writerow([
                    l.lamay_id,
                    venue_str,
                    l.deceased.full_name,
                    f"{l.buyer.full_name} ({l.buyer.contact_number})",
                    f"{l.casket.casket_id} - {l.casket.model}" if l.casket else 'N/A',
                    f"{l.lamay_start_date} {l.lamay_start_time or ''}".strip(),
                    str(l.expected_end_date or ''),
                    burial_str,
                    l.get_status_display()
                ])

        elif module == 'chapel_occupancy':
            # Requirement 24: CHAPEL OCCUPANCY REPORT
            writer.writerow(['Chapel', 'Deceased', 'Start Date', 'End Date', 'Status', 'Casket', 'Buyer'])
            lamays = LamayRecord.objects.select_related('chapel', 'deceased', 'buyer', 'casket').all()
            if start_date:
                lamays = lamays.filter(lamay_start_date__gte=start_date)
            if end_date:
                lamays = lamays.filter(lamay_start_date__lte=end_date)
            chapel_id = request.query_params.get('chapel')
            if chapel_id:
                lamays = lamays.filter(chapel_id=chapel_id)
            for l in lamays:
                venue_str = l.chapel.name if l.chapel else (f"Residence: {l.residence_address}" if l.residence_address else 'Residence / Home Viewing')
                writer.writerow([
                    venue_str,
                    l.deceased.full_name,
                    f"{l.lamay_start_date} {l.lamay_start_time or ''}".strip(),
                    f"{l.expected_burial_date or l.expected_end_date or ''} {l.burial_time or ''}".strip(),
                    l.get_status_display(),
                    f"{l.casket.casket_id} - {l.casket.model}" if l.casket else 'N/A',
                    l.buyer.full_name
                ])

        elif module == 'maintenance':
            writer.writerow(['ID', 'Location', 'Category', 'Issue', 'Priority', 'Reported By', 'Date Reported', 'Assigned To', 'Status', 'Cost (PHP)', 'Completed Date'])
            reqs = Maintenance.objects.select_related('location').all()
            reqs = get_target_loc_filter(reqs, 'location_id')
            if start_date:
                reqs = reqs.filter(date_reported__gte=start_date)
            if end_date:
                reqs = reqs.filter(date_reported__lte=end_date)
            for m in reqs:
                writer.writerow([
                    m.maintenance_id, m.location.name, m.category, m.issue, m.priority,
                    m.reported_by, m.date_reported, m.assigned_to, m.status, m.cost, m.date_completed or ''
                ])

        elif module == 'water':
            writer.writerow(['Location', 'Provider', 'Meter Number', 'Billing Period', 'Bill Date', 'Due Date', 'Prev Reading', 'Current Reading', 'Consumption (m3)', 'Amount (PHP)', 'Status', 'Date Paid'])
            bills = WaterBill.objects.select_related('location').all()
            bills = get_target_loc_filter(bills, 'location_id')
            for b in bills:
                writer.writerow([
                    b.location.name, b.provider, b.meter_number, b.billing_period,
                    b.bill_date, b.due_date, b.previous_reading, b.current_reading,
                    b.consumption, b.amount, b.payment_status, b.date_paid or ''
                ])

        elif module == 'electricity':
            writer.writerow(['Location', 'Provider', 'Meter Number', 'Billing Period', 'Bill Date', 'Due Date', 'Prev Reading', 'Current Reading', 'Consumption (kWh)', 'Amount (PHP)', 'Status', 'Date Paid'])
            bills = ElectricityBill.objects.select_related('location').all()
            bills = get_target_loc_filter(bills, 'location_id')
            for b in bills:
                writer.writerow([
                    b.location.name, b.provider, b.meter_number, b.billing_period,
                    b.bill_date, b.due_date, b.previous_reading, b.current_reading,
                    b.consumption, b.amount, b.payment_status, b.date_paid or ''
                ])

        return response


class ExportExcelView(APIView):
    """
    Generates a real Excel (.xlsx) workbook using openpyxl.
    Accepts the same query parameters as ExportCSVView:
      module, location, start_date, end_date
    Plus module-specific filters (search, status, category, priority, type).
    Enforces the same CanViewReports permission and location isolation.
    """
    permission_classes = [CanViewReports]

    def get(self, request):
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment
        except ImportError:
            return HttpResponse(
                'openpyxl is not installed. Run: pip install openpyxl',
                status=500,
                content_type='text/plain'
            )

        user = request.user
        accessible_loc_ids = user.get_accessible_location_ids()
        module = request.query_params.get('module', 'inventory').lower()
        location_id = request.query_params.get('location')
        start_date = request.query_params.get('start_date')
        end_date = request.query_params.get('end_date')
        search = request.query_params.get('search', '').strip()
        status_param = request.query_params.get('status', '').strip()

        # Helper to apply location isolation
        def get_target_loc_filter(qs, loc_field='location_id'):
            if not user.is_master_admin:
                qs = qs.filter(**{f"{loc_field}__in": accessible_loc_ids})
            if location_id:
                try:
                    loc_int = int(location_id)
                    if user.is_master_admin or loc_int in accessible_loc_ids:
                        qs = qs.filter(**{loc_field: loc_int})
                    else:
                        return qs.none()
                except ValueError:
                    qs = qs.filter(**{loc_field: location_id})
            return qs

        wb = Workbook()
        ws = wb.active

        # Header style
        header_font = Font(bold=True, color='FFFFFF')
        header_fill = PatternFill(fill_type='solid', fgColor='1E3A5F')
        header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

        def write_header(ws, headers):
            ws.append(headers)
            for cell in ws[1]:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment

        def freeze_and_autofit(ws):
            ws.freeze_panes = 'A2'
            for col in ws.columns:
                max_len = 0
                col_letter = col[0].column_letter
                for cell in col:
                    try:
                        if cell.value:
                            max_len = max(max_len, len(str(cell.value)))
                    except Exception:
                        pass
                ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 50)

        filename = f'Alaala_{module.replace("_", " ").title().replace(" ", "_")}'

        if module == 'inventory':
            ws.title = 'Inventory'
            write_header(ws, ['ID', 'Item Name', 'Category', 'Location', 'Quantity', 'Unit', 'Min Stock', 'Cost (PHP)', 'Total Value (PHP)', 'Status'])
            items = InventoryItem.objects.select_related('category', 'location').filter(is_active=True)
            items = get_target_loc_filter(items, 'location_id')
            if search:
                items = items.filter(
                    Q(item_name__icontains=search) |
                    Q(description__icontains=search) |
                    Q(supplier__icontains=search)
                )
            if status_param:
                items = items.filter(status=status_param)
            category_param = request.query_params.get('category', '').strip()
            if category_param:
                items = items.filter(category_id=category_param)
            for it in items:
                val = it.current_quantity * float(it.cost)
                ws.append([it.id, it.item_name, it.category.name, it.location.name, it.current_quantity, it.unit, it.minimum_stock, float(it.cost), round(val, 2), it.status])
            filename = 'Alaala_Inventory'

        elif module == 'transactions':
            ws.title = 'Transactions'
            write_header(ws, ['Date', 'Type', 'Item', 'Qty', 'From Location', 'To Location', 'Reason', 'User', 'Notes'])
            txs = InventoryTransaction.objects.select_related('item', 'from_location', 'to_location', 'user').all()
            if not user.is_master_admin:
                txs = txs.filter(Q(from_location_id__in=accessible_loc_ids) | Q(to_location_id__in=accessible_loc_ids))
            if location_id:
                txs = txs.filter(Q(from_location_id=location_id) | Q(to_location_id=location_id))
            if start_date and start_date.strip():
                txs = txs.filter(created_at__date__gte=start_date.strip())
            if end_date and end_date.strip():
                txs = txs.filter(created_at__date__lte=end_date.strip())
            type_param = request.query_params.get('type', '').strip()
            if type_param:
                txs = txs.filter(transaction_type=type_param)
            for t in txs:
                ws.append([
                    t.created_at.strftime('%Y-%m-%d %H:%M'),
                    t.transaction_type,
                    t.item.item_name,
                    t.quantity,
                    t.from_location.name if t.from_location else '',
                    t.to_location.name if t.to_location else '',
                    t.reason,
                    t.user.username if t.user else '',
                    t.notes
                ])
            filename = 'Alaala_Inventory_Transactions'

        elif module == 'caskets':
            ws.title = 'Caskets'
            write_header(ws, ['Casket ID', 'Model', 'Type', 'Size', 'Color', 'Material', 'Location', 'Cost (PHP)', 'Price (PHP)', 'Condition', 'Status', 'Received Date'])
            caskets = Casket.objects.select_related('location').all()
            caskets = get_target_loc_filter(caskets, 'location_id')
            if search:
                caskets = caskets.filter(
                    Q(casket_id__icontains=search) | Q(model__icontains=search) |
                    Q(material__icontains=search) | Q(color__icontains=search)
                )
            if status_param:
                caskets = caskets.filter(status=status_param)
            condition_param = request.query_params.get('condition', '').strip()
            if condition_param:
                caskets = caskets.filter(condition=condition_param)
            for c in caskets:
                ws.append([
                    c.casket_id, c.model, c.casket_type, c.size, c.color, c.material,
                    c.location.name, float(c.purchase_cost), float(c.selling_price),
                    c.condition, c.status, str(c.date_received)
                ])
            filename = 'Alaala_Caskets'

        elif module == 'lamay':
            ws.title = 'Chapel & Lamay'
            write_header(ws, ['Lamay ID', 'Chapel / Venue', 'Deceased', 'Buyer / Family Contact', 'Casket', 'Lamay Start', 'Expected End', 'Burial Date', 'Status'])
            lamays = LamayRecord.objects.select_related('chapel', 'deceased', 'buyer', 'casket').all()
            if start_date:
                lamays = lamays.filter(lamay_start_date__gte=start_date)
            if end_date:
                lamays = lamays.filter(lamay_start_date__lte=end_date)
            chapel_id = request.query_params.get('chapel')
            if chapel_id:
                lamays = lamays.filter(chapel_id=chapel_id)
            if status_param:
                lamays = lamays.filter(status=status_param)
            if search:
                lamays = lamays.filter(
                    Q(lamay_id__icontains=search) |
                    Q(deceased__full_name__icontains=search) |
                    Q(buyer__full_name__icontains=search)
                )
            for l in lamays:
                burial_str = f"{l.expected_burial_date} {l.burial_time or ''}".strip() if l.expected_burial_date else 'TBD'
                venue_str = l.chapel.name if l.chapel else (f"Residence: {l.residence_address}" if l.residence_address else 'Residence / Home Viewing')
                ws.append([
                    l.lamay_id,
                    venue_str,
                    l.deceased.full_name,
                    f"{l.buyer.full_name} ({l.buyer.contact_number})",
                    f"{l.casket.casket_id} - {l.casket.model}" if l.casket else 'N/A',
                    f"{l.lamay_start_date} {l.lamay_start_time or ''}".strip(),
                    str(l.expected_end_date or ''),
                    burial_str,
                    l.get_status_display()
                ])
            filename = 'Alaala_Chapel_Lamay'

        elif module == 'maintenance':
            ws.title = 'Maintenance'
            write_header(ws, ['ID', 'Location', 'Category', 'Issue', 'Priority', 'Reported By', 'Date Reported', 'Assigned To', 'Status', 'Cost (PHP)', 'Completed Date'])
            reqs = Maintenance.objects.select_related('location').all()
            reqs = get_target_loc_filter(reqs, 'location_id')
            if search:
                reqs = reqs.filter(
                    Q(maintenance_id__icontains=search) | Q(issue__icontains=search) |
                    Q(assigned_to__icontains=search) | Q(reported_by__icontains=search)
                )
            if status_param:
                reqs = reqs.filter(status=status_param)
            category_param = request.query_params.get('category', '').strip()
            if category_param:
                reqs = reqs.filter(category=category_param)
            priority_param = request.query_params.get('priority', '').strip()
            if priority_param:
                reqs = reqs.filter(priority=priority_param)
            if start_date:
                reqs = reqs.filter(date_reported__gte=start_date)
            if end_date:
                reqs = reqs.filter(date_reported__lte=end_date)
            for m in reqs:
                ws.append([
                    m.maintenance_id, m.location.name, m.category, m.issue, m.priority,
                    m.reported_by, str(m.date_reported), m.assigned_to or '',
                    m.status, float(m.cost), str(m.date_completed) if m.date_completed else ''
                ])
            filename = 'Alaala_Maintenance'

        elif module == 'water':
            ws.title = 'Water Bills'
            write_header(ws, ['Location', 'Provider', 'Meter Number', 'Billing Period', 'Bill Date', 'Due Date', 'Prev Reading', 'Current Reading', 'Consumption (m³)', 'Amount (PHP)', 'Status', 'Date Paid'])
            bills = WaterBill.objects.select_related('location').all()
            bills = get_target_loc_filter(bills, 'location_id')
            if search:
                bills = bills.filter(
                    Q(meter_number__icontains=search) | Q(billing_period__icontains=search) |
                    Q(provider__icontains=search)
                )
            if status_param:
                bills = bills.filter(payment_status=status_param)
            for b in bills:
                ws.append([
                    b.location.name, b.provider, b.meter_number, b.billing_period,
                    str(b.bill_date), str(b.due_date), float(b.previous_reading),
                    float(b.current_reading), float(b.consumption), float(b.amount),
                    b.payment_status, str(b.date_paid) if b.date_paid else ''
                ])
            filename = 'Alaala_Water'

        elif module == 'electricity':
            ws.title = 'Electricity Bills'
            write_header(ws, ['Location', 'Provider', 'Meter Number', 'Billing Period', 'Bill Date', 'Due Date', 'Prev Reading', 'Current Reading', 'Consumption (kWh)', 'Amount (PHP)', 'Status', 'Date Paid'])
            bills = ElectricityBill.objects.select_related('location').all()
            bills = get_target_loc_filter(bills, 'location_id')
            if search:
                bills = bills.filter(
                    Q(meter_number__icontains=search) | Q(billing_period__icontains=search) |
                    Q(provider__icontains=search)
                )
            if status_param:
                bills = bills.filter(payment_status=status_param)
            for b in bills:
                ws.append([
                    b.location.name, b.provider, b.meter_number, b.billing_period,
                    str(b.bill_date), str(b.due_date), float(b.previous_reading),
                    float(b.current_reading), float(b.consumption), float(b.amount),
                    b.payment_status, str(b.date_paid) if b.date_paid else ''
                ])
            filename = 'Alaala_Electricity'

        else:
            ws.title = 'Export'
            ws.append(['No data for unknown module:', module])

        freeze_and_autofit(ws)

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        response = HttpResponse(
            buffer.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}.xlsx"'
        return response

