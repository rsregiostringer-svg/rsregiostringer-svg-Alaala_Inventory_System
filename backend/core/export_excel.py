import openpyxl
from django.http import HttpResponse
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from utilities.models import WaterBill
from django.db.models import Q
from rest_framework.views import APIView
from rest_framework import permissions

class ExportXLSXView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        module = request.query_params.get('module')
        
        if module == 'water':
            location_id = request.query_params.get('location')
            water_source = request.query_params.get('water_source')
            search = request.query_params.get('search')
            date = request.query_params.get('date')

            qs = WaterBill.objects.select_related('location').all()

            if location_id:
                qs = qs.filter(location_id=location_id)
            if water_source:
                qs = qs.filter(water_source=water_source)
            if date:
                qs = qs.filter(date=date)
            if search:
                qs = qs.filter(
                    Q(tank__icontains=search) |
                    Q(patient_name__icontains=search) |
                    Q(shift__icontains=search)
                )

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Water Monitoring"

            # Define headers
            headers = [
                "Date", "Tenant / Location", "Name of Patient", "Shift", "Tank No.",
                "Water Level (In / Going)", "Additional Water", "Checked By", "Remarks",
                "Water Level (Out / Going)", "Additional Water", "Checked By", "Remarks",
                "Total Consumed", "No. of Refilled Gallons", "Encoded By", "Water Source"
            ]
            
            # 2-row header to handle merged cells as requested
            ws.append([""] * 5 + ["Initial Reading / Water In"] * 4 + ["Subsequent Reading / Water Out"] * 4 + [""] * 4)
            ws.append(headers)

            # Merge headers
            ws.merge_cells(start_row=1, start_column=6, end_row=1, end_column=9)
            ws.merge_cells(start_row=1, start_column=10, end_row=1, end_column=13)
            
            # Styling headers
            header_font = Font(bold=True)
            header_fill_in = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")
            header_fill_out = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
            alignment = Alignment(horizontal="center", vertical="center")
            
            for cell in ws[1]:
                cell.font = header_font
                cell.alignment = alignment
                if 6 <= cell.column <= 9:
                    cell.fill = header_fill_in
                elif 10 <= cell.column <= 13:
                    cell.fill = header_fill_out

            for cell in ws[2]:
                cell.font = header_font
                cell.alignment = alignment

            for row in qs:
                ws.append([
                    str(row.date) if row.date else '',
                    row.location.name if row.location else '',
                    row.patient_name or '',
                    row.shift or '',
                    row.tank or '',
                    float(row.initial_level) if row.initial_level is not None else '',
                    float(row.initial_additional) if row.initial_additional is not None else '',
                    row.initial_checked_by or '',
                    row.initial_remarks or '',
                    float(row.subsequent_level) if row.subsequent_level is not None else '',
                    float(row.subsequent_additional) if row.subsequent_additional is not None else '',
                    row.subsequent_checked_by or '',
                    row.subsequent_remarks or '',
                    float(row.consumption) if row.consumption is not None else '',
                    float(row.refilled_gallons) if row.refilled_gallons is not None else '',
                    row.encoded_by_name or '',
                    row.water_source or ''
                ])

            response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
            response['Content-Disposition'] = 'attachment; filename="Water_Monitoring.xlsx"'
            wb.save(response)
            return response
            
        return HttpResponse(status=400)
